# Instalar a biblioteca openpyxl para manipulação de Excel, caso não esteja presente
!pip install openpyxl -q

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from google.colab import files

# 1. Dados de exemplo estruturados (ou substitua pelo seu DataFrame do Supabase)
dados = [
    {"ID do Registo": 1791936063076, "Data e Hora": "2026-10-14 00:01:03", "Operador": "Analista SOC 01", "Nível de Severidade": "CRÍTICO", "Termo Técnico / Ameaça": "zero-day vulnerability in gateway"},
    {"ID do Registo": 1791841701321, "Data e Hora": "2026-10-12 21:48:21", "Operador": "Analista de teste 1", "Nível de Severidade": "Normal", "Termo Técnico / Ameaça": "zero trust network access architecture"},
    {"ID do Registo": 1791841662026, "Data e Hora": "2026-10-12 21:47:42", "Operador": "Analista de teste 1", "Nível de Severidade": "CRÍTICO", "Termo Técnico / Ameaça": "ransomware encryption bypass techniques"}
]

df = pd.DataFrame(dados)

# 2. Criar a estrutura do Excel com openpyxl
wb = Workbook()
ws = wb.active
ws.title = "Relatório Executivo"

# Adicionar cabeçalho
cabecalhos = list(df.columns)
ws.append(cabecalhos)

# Adicionar dados linha a linha
for _, linha in df.iterrows():
    ws.append(list(linha))

# 3. Estilização corporativa e formatação visual
fill_cabecalho = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Azul corporativo
font_cabecalho = Font(name="Arial", size=11, bold=True, color="FFFFFF")

borda_fina = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

fill_critico = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Vermelho suave para críticos
font_critico = Font(name="Arial", size=10, color="C00000", bold=True)

fill_normal = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Verde suave para normais
font_normal = Font(name="Arial", size=10, color="375623")

# Formatar o cabeçalho
for col_num in range(1, len(cabecalhos) + 1):
    celula = ws.cell(row=1, column=col_num)
    celula.fill = fill_cabecalho
    celula.font = font_cabecalho
    celula.alignment = Alignment(horizontal="center", vertical="center")

# Formatar os dados e aplicar cores condicionais
for row_num in range(2, len(df) + 2):
    severidade_val = ws.cell(row=row_num, column=4).value # Coluna Nível de Severidade
    is_critico = "CRÍTICO" in str(severidade_val)
    
    for col_num in range(1, len(cabecalhos) + 1):
        celula = ws.cell(row=row_num, column=col_num)
        celula.border = borda_fina
        celula.alignment = Alignment(vertical="center")
        
        if is_critico:
            celula.fill = fill_critico
            if col_num == 4:
                celula.font = font_critico
        else:
            if col_num == 4:
                celula.fill = fill_normal
                celula.font = font_normal

# 4. Ajustar larguras das colunas automaticamente
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for celula in col:
        if celula.value:
            max_len = max(max_len, len(str(celula.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

# 5. Guardar o ficheiro e iniciar o download imediato no browser
nome_ficheiro = "relatorio_ciber_estruturado.xlsx"
wb.save(nome_ficheiro)

print(f"Ficheiro '{nome_ficheiro}' gerado com sucesso!")
files.download(nome_ficheiro)
